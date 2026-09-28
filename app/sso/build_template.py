# -*- coding: utf-8 -*-
"""
SSO_Prerequis_Template_EN.xlsx / _FR.xlsx : le fichier SSO envoyé au client après la réunion de lancement.
Reprend la trame Equadis (Task List / Owner / Value [Production] / Value [Test] / What it implies),
avec « Client_Name » à la place du nom du client.
"""
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F = 'Calibri'
NAVY = '1F2A36'
GREEN = '00E05A'
RED = 'FF0000'
side = Side(style='thin', color='000000')
box = Border(left=side, right=side, top=side, bottom=side)
wrap = Alignment(wrap_text=True, vertical='center')


TXT = {
 'EN': dict(task='Task List', owner='Owner', prod='Value [Production]', test='Value [Test]', impl='What it implies',
   app='Application Name', signin='Application Sign in & Sign out URL or Internal/External URL (optional)',
   redirect='Redirect URL', name='NAME', lst='(optional) List of Client_Name users who will access the application',
   lstv='List user of Client_Name', i_app='User friendly name for the application',
   i_sign='URL with which users will sign-in to the application. Optional if app supports IDP initiated SSO',
   i_red='URL where the application receive the authentication token'),
 'FR': dict(task='Liste des tâches', owner='Responsable', prod='Valeur [Production]', test='Valeur [Test]', impl='Ce que cela implique',
   app="Nom de l'application", signin="URL de connexion et de déconnexion de l'application, ou URL interne/externe (optionnel)",
   redirect='URL de redirection', name='NOM', lst="(optionnel) Liste des utilisateurs Client_Name qui accéderont à l'application",
   lstv='Liste des utilisateurs Client_Name', i_app="Nom de l'application affiché aux utilisateurs",
   i_sign="URL avec laquelle les utilisateurs se connectent à l'application. Optionnel si l'application prend en charge le SSO initié par le fournisseur d'identité",
   i_red="URL où l'application reçoit le jeton d'authentification"),
}

def construire(lang):
    t = TXT[lang]
    wb = Workbook(); ws = wb.active; ws.title = 'SSO'
    for col, w in zip('ABCDE', (38, 18, 42, 42, 36)):
        ws.column_dimensions[col].width = w

    # en-tête : logo + titre
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 30
    ws.merge_cells('A1:A2'); ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    logo = Image('equadis_logo_white.png'); logo.width, logo.height = 216, 53
    ws.add_image(logo, 'A1')
    ws.merge_cells('B1:D2'); ws['B1'] = 'SSO'
    ws['B1'].font = Font(name=F, size=20, color=NAVY); ws['B1'].alignment = Alignment(horizontal='left', vertical='center', indent=3)
    ws.merge_cells('E1:E2')
    for r in (1, 2):
        for c in 'ABCDE': ws[f'{c}{r}'].border = box

    # ligne de titres
    ws.row_dimensions[3].height = 34
    for c, titre, coul in [('A', t['task'], 'FFFFFF'), ('B', t['owner'], 'FFFFFF'), ('C', t['prod'], GREEN),
                       ('D', t['test'], GREEN), ('E', t['impl'], 'FFFFFF')]:
        x = ws[f'{c}3']; x.value = titre; x.fill = PatternFill('solid', fgColor=NAVY)
        x.font = Font(name=F, size=11, bold=True, color=coul); x.alignment = Alignment(horizontal='center', vertical='center'); x.border = box

    gras = Font(name=F, size=11, bold=True)
    lien = Font(name=F, size=11, color='0563C1', underline='single')
    rouge = Font(name=F, size=11, bold=True, color=RED)

    def ligne(r, tache, owner, prod, test, implique, owner_font=gras, val_font=None, h=None):
        for c, v in zip('ABCDE', (tache, owner, prod, test, implique)):
            x = ws[f'{c}{r}']; x.value = v; x.border = box; x.alignment = wrap; x.font = gras
        ws[f'B{r}'].font = owner_font
        if val_font:
            ws[f'C{r}'].font = val_font; ws[f'D{r}'].font = val_font
        if h: ws.row_dimensions[r].height = h

    ligne(4, t['app'], 'EQUADIS', 'Gaia', 'Gaia', t['i_app'])
    ligne(5, t['signin'], 'EQUADIS',
          'https://gaia.equadis.com/', 'https://customer.equadis.com/',
          t['i_sign'], val_font=lien, h=60)
    ligne(6, None, None, None, None, None)
    ws.merge_cells('A7:A8'); ws.merge_cells('B7:B8'); ws.merge_cells('C7:C8'); ws.merge_cells('D7:D8'); ws.merge_cells('E7:E8')
    ligne(7, t['redirect'], 'EQUADIS Génération URLs PAR TEAM GAIA',
          'https://gaia.equadis.com/#login-sso/Client_Name', 'https://customer.equadis.com/#login-sso/Client_Name',
          t['i_red'], val_font=lien, h=30)
    for c in 'ABCDE': ws[f'{c}8'].border = box
    ws.row_dimensions[8].height = 30
    ligne(9, t['name'], None, None, None, None)
    for r, champ in enumerate(['Application ID', 'Secret Key', 'Discovery URL', '(If not discovery URL issuer) issuer',
                           '(If not discovery URL issuer) token_url', '(If not discovery URL issuer) scope',
                           '(If not discovery URL issuer) userinfo_url', '(optional) Redirect URL'], 10):
        ligne(r, champ, 'Client_Name', None, None, None, owner_font=rouge)
    ligne(18, t['lst'], 'EQUADIS',
          t['lstv'], t['lstv'], None, val_font=Font(name=F, size=11, bold=True, color=RED), h=34)

    ws.sheet_view.showGridLines = True
    wb.save(f'SSO_Prerequis_Template_{lang}.xlsx')

for lang in ('EN', 'FR'):
    construire(lang)
print('ok')
