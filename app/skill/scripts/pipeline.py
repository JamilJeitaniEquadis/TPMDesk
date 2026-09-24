# -*- coding: utf-8 -*-
"""
Chaîne de traitement générique : export SupplierXM -> fichier au format de l'extraction Gaia.
Les colonnes source sont repérées par leur chemin technique (ligne 4), pas par leur lettre :
Salsify n'exporte que les colonnes servies, donc la position varie d'un export à l'autre.
"""
import re, os, json, shutil, zipfile, datetime
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as gl

import argparse, os as _os
TPL = None          # renseigné par la ligne de commande
FIRST_NEW, LAST_TPL = 7, 16
_ICI = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))

# ============================ tables de valeurs ============================
V_PACK = {'BO':'Flacon','BX':'Boîte','CT':'Carton','TU':'Tube','JR':'Pot (GDSN)','NE':'Non emballé',
          'CY':'Cylindre','BBG':'Bag in Box','WRP':'Emballage sous pellicule','STL':'Stick','PUG':'Colis',
          'EN':'Enveloppe (EN)','CKT':'Boîte','BPG':'Emballage sous blister (BPG)',
          'PX':'Palette (PX)','SX':'Multipack (MPG)','TN':'Boîte (CNG)','CNG':'Boîte (CNG)',
          'AM':'Ampoule','AE':'Aérosol','CU':'Gobelet / Pot / Bol (CU)','PO':'Poche','PT':'Pot (GS1 France)',
          'CR':'Caisse / casier à bouteilles','CS':'Caisse isotherme','BG':'Gros sac / dimensions palette',
          'SH':None}
V_GENDER = {'female':'Femme','male':'Homme','mixed':'Unisexe','unisex':'Unisexe',
            'unclassified':'Unisexe','unidentified':'Unisexe'}
V_BARCODE = {'EAN_UCC_13_SYMBOL':'EAN 13','EAN_UCC_8_SYMBOL':'EAN 8','ITF_14':'ITF 14',
             'ITF_14_SYMBOL':'ITF 14','NONE':'Aucun code barre'}
V_ORDER_UOM = {'pce':'Pièce','L':'Litre','cnt':'UVC','kg':'Kilogramme','g':'Gramme','mL':'Millilitre'}
V_COLOR = {'UNSPECIFIED':'Indéterminé','Brown':'Marron Moyen','Pink':'Rose Moyen','Purple':'Violet Moyen',
           'Blue':'Bleu Moyen','Black':'Noir','White':'Blanc','LIGHT_BEIGE':'Beige Clair','Nude':'Beige Clair',
           'LIGHT_BLUE':'Bleu Clair','LIGHT_PINK':'Rose Clair','Pinkish':'Rose Clair','Orange':'Orange Moyen',
           'Red':'Rouge Moyen','Yellow':'Jaune Moyen','Green':'Vert Moyen','Grey':'Gris Moyen','Gray':'Gris Moyen',
           'Beige':'Beige Moyen','Gold':'Doré','GOLD':'Doré','GOLDEN':'Doré','Silver':'Argenté','SILVER':'Argenté','Clear':'Incolore','Khaki':'Kaki Moyen',
           'DARK_BEIGE':'Beige Foncé','DARK_BROWN':'Marron Foncé','LIGHT_BROWN':'Marron Clair','DARK_BLUE':'Bleu Foncé',
           'DARK_PINK':'Rose Foncé','LIGHT_RED':'Rouge Clair','DARK_RED':'Rouge Foncé','COLOURLESS':'Incolore',
           'Transparent':'Incolore','LIGHT_PURPLE':'Violet Clair','DARK_PURPLE':'Violet Foncé','LIGHT_GREEN':'Vert Clair',
           'DARK_GREEN':'Vert Foncé','LIGHT_YELLOW':'Jaune Clair','DARK_YELLOW':'Jaune Foncé','LIGHT_GREY':'Gris Clair',
           'DARK_GREY':'Gris Foncé','LIGHT_ORANGE':'Orange Clair','DARK_ORANGE':'Orange Foncé'}
V_DANGER = {'3':'Danger de feu (matière liquide inflammable)','4.1':'Danger de feu (matière solide inflammable)',
            '2':'Gaz','2.1':'Gaz inflammable et non toxique','2.2':'Gaz non inflammable et non toxique',
            '2.3':'Gaz toxique','5.1':'Matières comburantes','5.2':"Peroxyde organique Danger d'incendie",
            '6.1':'Matière toxique','6.2':'Matière infectieuse','8':'Matière corrosive',
            '9':'Matières et objets divers présentant, au cours du transport, un danger autre que ceux visés par les autres classes'}
V_SUBST = {'alco':'ALCOHOL (ALCO)','std':'STANDARD (STD)','aero':'AEROSOL (AERO)','solv':'SOLVENT (SOLV)'}
V_SPF = {'1':'6 (Faible)','2':'10 (Faible)','3':'15 (Moyenne)','4':'20 (Moyenne)','5':'25 (Moyenne)',
         '6':'30 (Haute)','7':'50 (Haute)','8':'50+ (Très Haute)'}
V_AGE = {'MIDDLE_ADULT':'Adulte','UNCLASSIFIED':'Non classé','ALL_AGES':'Tout âge','CHILD':'Enfant',
         'BABY':'Bébé','TEENAGER':'Adolescent','SENIOR':'Senior','NEWBORN':'Nouveau né'}
V_SKIN = {'ALL_SKINS':'Tous types de peaux (gs1 standard)','ALL_SKIN_TYPES':'Tous types de peaux (gs1 standard)',
          'COMBINATION_SKIN':'mixte (gs1 standard)','NORMAL_SKIN':'normale (gs1 standard)',
          'OILY_SKIN':'grasse (gs1 standard)','MATURE_SKIN':'mature (gs1 standard)',
          'VERY_DRY_SKIN':'Très Sèche (gs1 standard)','DRY_SKIN':'sèche (gs1 standard)',
          'SENSITIVE_SKIN':'sensible (gs1 standard)','ACNE_PRONE_SKIN':'acnéique (gs1 standard)',
          'MIXED_TO_OILY_SKIN':'mixte (gs1 standard)','DRY_TO_DEHYDRATED_SKIN':'sèche (gs1 standard)',
          'YOUNG_SKIN':'Peau Jeune','ATOPIC_SKIN':'Peau atopique','TIRED_SKIN':'Peau fatiguée',
          'OTHER':'Autre'}
V_TEXTURE = {'LIQUID':'Liquide','CREAM':'Crème','MIST':'Eau / Brume','WATER':'Eau / Brume','OTHER':'Autre',
             'FLUID':'Fluide','POWDER':'Poudre','SERUM':'Serum','SPRAY':'Spray','GEL':'Gel','OIL':'Huile',
             'BALM':'Baume','PRESSED_POWDER':'Compacte','COMPACT':'Compacte','STICK':'Stick','MILK':'Lait',
             'MILK_COSMETICS':'Lait','FOAM':'Mousse','WAX':'Cire','PASTE':'Pate','SOLID':'Solide','SOLIDE':'Solide',
             'LOOSE_POWDER':'Poudre Libre','LOTION':'Lotion','WIPE':'Lingette','PATCH':'Patch','BIPHASE':'Bi-phase',
             'EXFOLIATING':'Exfoliante','EXFOLIANT':'Exfoliante','SUPPLE':'Souple','CREAMY':'Crémeux','RIGID':'Rigide',
             'BUTTER_COSMETICS':'Crémeux','SOAP':'Solide','CAPSULE':'Dosette','GRANULES':'Granules',
             'SHEET':'Tissus','CRYSTAL':'Solide','MASK_ABSORB':'Patch','CONCENTRATE':'Serum',
             'DRY_SHAMPOO':'Shampoing Sec','EFFERVESCENT':'Effervescent','WIPES':'Lingette'}
V_AREA = {'BODY':'Corps','BODY_SKIN':'Corps','FACE':'Visage','NECK':'Cou','EYES':'Yeux','LIPS':'Lèvres',
          'HAIR':'Cheveux','NAILS':'Ongles','HAND':'Main','FEET':'Pieds','SKIN':'Peau','EYELASHES':'Cils',
          'EYEBROWS':'Sourcils','TEETH':'Dents','MOUTH':'Bouche','LEGS':'Jambes','ARM':'Bras','BEARD':'Barbe',
          'HEAD':'Tête','ARMPIT':'Aisselle','DECOLLETE':'Décolleté','NECKLINE':'Décolleté','CHEST':'Poitrine',
          'EYE_LASH':'Cils','EYEBROW':'Sourcils','EYE_BROWS':'Sourcils','ELBOW':'Coude','MULTI':'Peau',
          'OTHER':'Peau','INTIMATE':'Partie intime','BACK':'Dos','KNEE':'Genou','SHOULDER':'Épaule'}
V_CONTACT = {'MANUFACTURER':'Exploitant du secteur (BZL)','CONSUMER_SUPPORT':'SUPPORT CONSOMMATEUR (CXC)',
             'TARGET_MARKET_INFORMATION_PROVIDER':'Contact Information (IC)','WLS':'Grossiste (WLS)'}
V_CHANNEL = {'TELEPHONE':'Téléphone','EMAIL':'Email','WEBSITE':'Website','FAX':'Telefax','MOBILE':'Site mobile'}
V_COUNTRY = {'250':'France','410':'Corée; République','380':'Italie','276':'Allemagne','724':'Espagne',
             '826':"Royaume-Uni de Grande-Bretagne et d'Irlande du Nord",'840':'Etats-Unis','756':'Suisse',
             '056':'Belgique','528':'Pays-Bas','616':'Pologne','392':'Japon','156':'Chine','203':'Tchéquie',
             '620':'Portugal','372':'Irlande','040':'Autriche','752':'Suède','208':'Danemark','246':'Finlande',
             '792':'Turquie','504':'Maroc','788':'Tunisie','484':'Mexique','076':'Brésil','124':'Canada',
             '036':'Australie','554':'Nouvelle-Zélande','710':"Afrique du Sud (République d')",'643':'Russie',
             '348':'Hongrie','642':'Roumanie','300':'Grèce','705':'Slovénie','703':'Slovaquie','100':'Bulgarie',
             '191':'Croatie','442':'Luxembourg','578':'Norvège','352':'Islande','008':'Albanie','012':'Algérie',
             '360':'Indonésie','356':'Inde','764':'Thaïlande','704':'Viêt-nam','458':'Malaisie','702':'Singapour',
             '344':'Hong Kong','158':'Taiwan','032':'Argentine','152':'Chili','170':'Colombie','604':'Pérou'}
U_WEIGHT = {'g':'Gramme (g)','kg':'Kilogramme (kg)','mg':'Milligramme (mg)'}
U_DIM = {'cm':'Centimètre (cm)','mm':'Millimètre (mm)'}
U_CONTENT = {'mL':'Millilitre (ml)','L':'Litre (l)','g':'Gramme (g)','kg':'Kilogramme (kg)','pce':'Pièce'}
V_LOG_TYPE = {'CASE':'Carton','PALLET':'Palette','PACK':'Emballage intermédiaire'}
V_PLATFORM_TC = {'one-way_pallet':'Non réutilisable (perdu)','exchange_pallets':'Echangeable au point de livraison',
                 'no_exchange':'Support palette non écheangeable, non retournable',
                 'return_pallets':'Support palette retournable',
                 'returnable_pallet':"Support palette consigné. Le support palette doit être retourné au point d'expédition",
                 'third-party-exchange_pallet':'Loué'}
V_UOM_FACT = {'pce':'UVC','cnt':'UVC','L':'Litre (l)','kg':'Kilogramme (kg)'}

GPC = {'71542':'10000356','100817':'10000731','70050':'10000381','70053':'10000356',
       '70051':'10000345','101005':'10000535','70068':'10000365','70056':'10000356','70055':'10000356','70062':'10000533','70061':'10000532',
       '70063':'10000534','70054':'10000332','70058':'10000373','71551':'10000360','70049':'10000368',
       '100993':'10000330','100812':'10000338','70048':'10000346','70067':'10008024','100992':'10000330',
       '70059':'10000356'}
CLASSIF = {
 ('10000672','Homme'):"Cosmétiques | Parfum | Parfum Homme | Coffret parfum homme | 10000672 - Assortiments de Produits Cosmétiques/Parfums",
 ('10000672','Femme'):"Cosmétiques | Parfum | Parfum Femme | Coffret parfum femme | 10000672 - Assortiments de Produits Cosmétiques/Parfums",
 ('10000672','Unisexe'):"Cosmétiques | Parfum | Parfum mixte | Coffret parfum mixte | 10000672 - Assortiments de Produits Cosmétiques/Parfums",
 ('10000669',None):"Cosmétiques | Maquillage | Assortiments de Produits Cosmétiques/de Maquillage | Assortiments de Produits Cosmétiques/de Maquillage | 10000669 - Assortiments de Produits Cosmétiques",
 ('10000365','Homme'):"Cosmétiques | Parfum | Parfum Homme | Eau de parfum homme | 10000365 - Parfums",
 ('10000365','Femme'):"Cosmétiques | Parfum | Parfum Femme | Eau de parfum femme | 10000365 - Parfums",
 ('10000365','Unisexe'):"Cosmétiques | Parfum | Parfum mixte | Eau de parfum mixte | 10000365 - Parfums",
 ('10000356',None):"Cosmétiques | Soin du Corps | Soin spécifique corps - Cosmétique | Crème hydratante | 10000356 - Produit de Soin de la Peau/Hydratant",
 ('10000373',None):"Cosmétiques | Soin du Corps | Solaire | Crème solaire | 10000373 - Produits de Protection du Soleil",
 ('10000532',None):"Cosmétiques | Maquillage | Teint  | Fond de teint | 10000532 - Produits Cosmétiques – Teint",
 ('10000533',None):"Cosmétiques | Maquillage | Yeux | Mascara | 10000533 - Produits Cosmétiques - Yeux",
 ('10000534',None):"Cosmétiques | Maquillage | Lèvres  | Rouge à lèvres | 10000534 - Produits Cosmétiques - Lèvres",
 ('10005727',None):"Cosmétiques | Maquillage | Lèvres  | Baume à lèvres | 10005727 - Baumes à lèvres",
 ('10000360',None):"Cosmétiques | Maquillage | Ongles  | Vernis à ongles | 10000360 - Produits Cosmétiques - Ongles",
 ('10000368',None):"Cosmétiques | Cheveux | Soin Cheveux | Shampooing | 10000368 - Cheveux - Shampooing",
 ('10000346',None):"Cosmétiques | Cheveux | Soin Cheveux | Après shampooing | 10000346 - Cheveux - Produits Conditionneurs/Traitants",
 ('10000338',None):"Cosmétiques | Soin du Corps | Soin spécifique corps - Cosmétique | Déodorant | 10000338 - Anti-transpirants/Déodorants",
 ('10000330',None):"Cosmétiques | Soin du Corps | Bain et douche | Gel douche | 10000330 - Nettoyage/Toilette/Savon - Personnel",
 ('10000332',None):"Cosmétiques | Soin visage | Nettoyant et démaquillant - cosmétique | Démaquillant Visage cosmétique | 10000332 - Nettoyants/Démaquillants (Non Electriques)",
 ('10000381',None):"Cosmétiques | Cheveux | Coiffant et fixant | Gel et mousse | 10000381 - Cheveux - Produits Coiffants (Non Electriques)",
 ('10000345',None):"Cosmétiques | Cheveux | Soin Cheveux | Coloration cheveux | 10000345 - Cheveux - Produits pour la Coloration",
 ('10000535',None):"Cosmétiques | Soin visage | Rasage | Gel et mousse à raser | 10000535 - Produits de Préparation au Rasage",
 ('10000731',None):"Cosmétiques | Soin du Corps | Solaire | Auto bronzant Corps | 10000731 - Produits Bronzants - Oraux (Non Electriques)",
 ('10008024',None):"Cosmétiques | Accessoires | Pinceau ou éponge  | Pinceau, houpette ou éponge teint | 10008024 - Applicateurs de produits cosmétiques",
}

# ============================ utilitaires ============================
# chemins liés à une catégorie source précise, prioritaires sur la correspondance par brique
CLASSIF_KIND = {
 '70017': "Cosmétiques | Parfum | Parfum enfant | Eau de senteur | 10000365 - Parfums",
 '70052': "Cosmétiques | Accessoires | Cheveux - Accessoires | Brosse et peigne cheveux | 10008023 - Cheveux - Brosses et peignes",
 '70057': "Cosmétiques | Accessoires | Accessoires Cosmétiques/de Maquillage | Accessoires et Outils | 10000377 - Accessoires/Outils de Cosmétique",
 '71542': "Cosmétiques | Soin visage | Hydratant et nourrissant  | Contour des yeux hydratant et nourrissant  | 10000356 - Produit de Soin de la Peau/Hydratant",
}

def code(v):
    if v in (None, ''): return None
    s = str(v); return s.split('∣')[-1].strip() if '∣' in s else s.strip()
def lab(v):
    if v in (None, ''): return None
    s = str(v); return s.split('∣')[0].strip() if '∣' in s else s.strip()
def gtin_txt(v):
    return None if v in (None, '') else str(v).strip()
def d_fr(v):
    if v in (None, ''): return None
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%d/%m/%Y')
    s = str(v)[:10]
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})$', s)
    return f"{m.group(3)}/{m.group(2)}/{m.group(1)}" if m else s
def trunc(v, n):
    if v in (None, ''): return None
    s = str(v).strip()
    return s if len(s) <= n else (s[:n].rsplit(' ', 1)[0] if ' ' in s[:n] else s[:n])
def boo(v):
    c = (code(v) or '').lower()
    return 'Oui' if c in ('vrai','true','oui') else ('Non' if c in ('faux','false','non') else None)
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

class Feuille:
    """Accès aux colonnes d'une feuille d'export par chemin technique."""
    def __init__(self, ws):
        rows = list(ws.iter_rows(min_row=1, max_row=7, values_only=True))
        self.idx = defaultdict(list)
        for c in range(1, len(rows[0])):
            p = rows[3][c]
            if p: self.idx[str(p)].append((c, code(rows[5][c])))
        self.rows = list(ws.iter_rows(min_row=8, values_only=True))
    def v(self, row, path, occ=0):
        cols = self.idx.get(path, [])
        return row[cols[occ][0]] if occ < len(cols) else None
    def n(self, path):
        return len(self.idx.get(path, []))
    def vu(self, row, path):
        """retourne (valeur, unité) pour la première unité servie"""
        for c, u in self.idx.get(path, []):
            if row[c] not in (None, ''): return row[c], u
        return None, None

# ============================ construction des enregistrements ============================
def construire(export_path, log):
    wb = load_workbook(export_path, read_only=True, data_only=True)
    P, L = Feuille(wb['Product']), Feuille(wb['Logistical units'])

    # hiérarchie logistique
    bygtin = {str(L.v(r, 'gtin')): r for r in L.rows}
    parent = defaultdict(list)
    for r in L.rows:
        ch = L.v(r, 'children.gtin')
        if ch: parent[str(ch)].append(str(L.v(r, 'gtin')))
    def chaine(g):
        out, cur, seen = [], g, set()
        while len(out) < 3:
            ps = parent.get(cur, [])
            if not ps or cur in seen: break
            if len(ps) > 1:
                log.append({'gtin': g, 'champ': 'Hiérarchie logistique',
                            'constat': f"{len(ps)} parents pour {cur} : le premier est retenu."})
            seen.add(cur); cur = ps[0]; out.append(cur)
        return out

    prods, logis = [], []
    for r in P.rows:
        g = gtin_txt(P.v(r, 'gtin'))
        d, dl = {}, {}
        def S(fid, val, occ=0):
            if val not in (None, ''): d[(fid, occ)] = val
        
        def SL(fid, val, occ=0):
            if val not in (None, ''): dl[(fid, occ)] = val

        S('2', g); ref_int = P.v(r, 'supplierCode')
        if ref_int and len(str(ref_int)) > 24:
            log.append({'gtin': g, 'champ': 'Référence interne (3)',
                        'constat': f"Référence de {len(str(ref_int))} caractères (limite indicative 24) : écrite en entier."})
        S('3', ref_int); S('1', 'UVC'); S('41', 'Finale')
        S('8_1', 'France'); S('8_2', 'GSS Parfumerie')
        S('109', trunc(P.v(r, 'brandText'), 40)); S('112', trunc(P.v(r, 'subBrandText'), 70))
        nom = P.v(r, 'namePublicLong')
        S('97', trunc(nom, 35)); S('103', trunc(nom, 200))
        S('98', trunc(P.v(r, 'functionalName'), 500)); S('102', trunc(P.v(r, 'nameLegal'), 500))
        S('101', trunc(P.v(r, 'productRangeText'), 500))

        # classification
        k = code(P.v(r, 'kind')) or ''
        cls = (lab(P.v(r, 'isClassifiedIn')) or '').upper()
        web = lab(P.v(r, 'org1337WebsiteProductCategory')) or ''
        genre = V_GENDER.get(code(P.v(r, 'targetConsumerGender')))
        if 'COFFRET' in cls or 'Coffret' in web or 'coffret' in str(nom or '').lower():
            brick = '10000672' if (k == '70068' or 'ALCOOL' in cls) else '10000669'
        elif 'BAUME LEVRES' in cls:
            brick = '10005727'
        else:
            brick = GPC.get(k)
        chemin = CLASSIF_KIND.get(k) or CLASSIF.get((brick, genre)) or CLASSIF.get((brick, None)) or CLASSIF.get((brick, 'Unisexe'))
        if chemin is None:
            log.append({'gtin': g, 'champ': 'Classification (11)',
                        'constat': f"Catégorie source « {P.v(r,'kind')} » : aucun chemin défini."})
        S('11', chemin)

        # dates
        dispo = d_fr(P.v(r, 'consumerFirstAvailabilityDateTime'))
        S('46', dispo or '30/06/2026'); S('2230', dispo)
        S('2234', d_fr(P.v(r, 'firstOrderDateTime'))); S('2231', d_fr(P.v(r, 'firstShipDateTime')))

        # mesures
        for path, umap, fv, fu in [('isSizedBy.depth', U_DIM, '117', '118'),
                                   ('isSizedBy.width', U_DIM, '132', '133'),
                                   ('isSizedBy.height', U_DIM, '121', '122'),
                                   ('grossWeight', U_WEIGHT, '156', '157'),
                                   ('netWeight', U_WEIGHT, '158', '159'),
                                   ('netContent', U_CONTENT, '126', '127')]:
            val, u = P.vu(r, path)
            if val not in (None, ''):
                S(fv, val); S(fu, umap.get(u))

        for fid, val in [('35','Oui'),('36','Oui'),('37','Non'),('38','Non'),('40','Non'),
                         ('44','Non'),('1649','Non'),('160','Produit Permanent')]:
            S(fid, val)
        S('39', 'Oui' if code(P.v(r, 'lifeCycle')) == 'PURCHASABLE' else 'Non')
        S('552', boo(P.v(r, 'isDangerousSubstance')))
        S('1709', V_COUNTRY.get(code(P.v(r, 'countryOfSettlement'))))
        S('73', V_ORDER_UOM.get(code(P.v(r, 'orderingUnitOfMeasure'))))
        S('169', gtin_txt(P.v(r, 'replaces.targetProduct.gtin')))
        S('2617', P.v(r, 'packagingInformationList.packagingLevelNumber'))
        for i in range(min(P.n('isComplementaryWith.targetProduct.gtin'), 18)):
            S('2580', gtin_txt(P.v(r, 'isComplementaryWith.targetProduct.gtin', i)), i)

        for i in range(P.n('partyInformationList.partyRoleCode')):
            if code(P.v(r, 'partyInformationList.partyRoleCode', i)) == 'BRAND_OWNER':
                S('172', gtin_txt(P.v(r, 'partyInformationList.partyGLN', i)))
                S('174', trunc(P.v(r, 'partyInformationList.partyNameText', i), 70))
                break
        else:
            log.append({'gtin': g, 'champ': 'Propriétaire de la marque (172/174)',
                        'constat': "Aucun bloc organisation avec le rôle BRAND_OWNER."})

        S('255', trunc(P.v(r, 'composition'), 5000))
        S('198', trunc(P.v(r, 'advices'), 1500))
        S('1086', trunc(P.v(r, 'warnings.text'), 1500))
        S('2238', trunc(P.v(r, 'description'), 2000))
        for i in range(min(P.n('productBenefits.text'), 2)):
            S('400', trunc(P.v(r, 'productBenefits.text', i), 250), i)
        S('311', P.v(r, 'ratioAlcohol'))
        S('113', V_COLOR.get(code(P.v(r, 'colorGroup'))))
        S('114', trunc(P.v(r, 'colorDescription'), 80))
        S('721', genre)
        if P.v(r, 'flashPointTemperature') not in (None, ''):
            S('1428', P.v(r, 'flashPointTemperature')); S('1429', 'Degré Celsius (°C)')
        icpe = lab(P.v(r, 'iCPEStorageCompatibilityCode'))
        if icpe:
            m = re.match(r'^(\d+)', icpe); S('1454', m.group(1) if m else icpe[:35])
        S('1634', V_DANGER.get(code(P.v(r, 'classOfDangerousGoods'))))
        S('1646', P.v(r, 'unitedNationsDangerousGoodsNumber'))
        S('1229', V_COUNTRY.get(code(P.v(r, 'originCountry'))))
        S('1025', V_PACK.get(code(P.v(r, 'packagingInformationList.packagingTypeCode'))))
        S('713', trunc(P.v(r, 'packagingInformationList.packagingTypeDescriptionTextList'), 200))
        S('1078', boo(P.v(r, 'packagingInformationList.isPackagingReturnable')))
        S('1076', boo(P.v(r, 'hasBatchNumber')))
        S('1544', V_BARCODE.get(code(P.v(r, 'dataCarrierTypeCode'))))
        S('1403', V_AREA.get(code(P.v(r, 'areaOfUseList.areaOfUseCode'))))
        S('687', V_SPF.get(code(P.v(r, 'sunburnProtectionFactor'))))
        S('2320', P.v(r, 'hexadecimalCode'))
        S('2321', V_SUBST.get(code(P.v(r, 'typeOfDangerousSubstanceOrArticle'))))
        S('2322', trunc(P.v(r, 'headNote'), 500)); S('2323', trunc(P.v(r, 'heartNote'), 500))
        S('2324', trunc(P.v(r, 'baseNote'), 500))
        S('2158', V_AGE.get(code(P.v(r, 'targetConsumerAgeList.targetConsumerAgeCode'))))
        for i in range(min(P.n('skinTypeList.skinTypeCode'), 10)):
            S('1402', V_SKIN.get(code(P.v(r, 'skinTypeList.skinTypeCode', i))), i)
        for i in range(min(P.n('productTextureList.productTextureCode'), 10)):
            S('1404', V_TEXTURE.get(code(P.v(r, 'productTextureList.productTextureCode', i))), i)
        if boo(P.v(r, 'priceLegalCommonUnitObligation')) == 'Oui':
            _, u = P.vu(r, 'netContent')
            S('1501', {'mL':'Prix au litre','L':'Prix au litre','g':'Prix au kilo','kg':'Prix au kilo',
                       'pce':'Prix à la pièce'}.get(u, 'Pas de prix par unité de mesure'))
        else:
            S('1501', 'Pas de prix par unité de mesure')
        if P.v(r, 'importEuropeanClassification') not in (None, ''):
            S('1224', 'Intrastat'); S('1225', P.v(r, 'importEuropeanClassification'))
        if P.v(r, 'catalogPrice') not in (None, ''):
            S('2261', P.v(r, 'catalogPrice')); S('2542', 'Euro'); S('2262', 1); S('2263', 'UVC')
            S('2264', d_fr(P.v(r, 'catalogPriceStartDate')))
        S('605', V_COUNTRY.get(code(P.v(r, 'dutyFeeTaxInformationList.dutyFeeTaxCountryCode'))))
        if code(P.v(r, 'dutyFeeTaxInformationList.dutyFeeTaxTypeCode')) == 'EXEMPT':
            S('604', 'Exoneré 0%')
        else:
            taux = num(P.v(r, 'dutyFeeTaxInformationList.dutyFeeTaxList.dutyFeeTaxRateNumber'))
            S('604', {20:'TVA 20 %', 5.5:'TAUX REDUIT 5.5 %', 10:'TAUX INTERMEDIAIRE 10%',
                      2.1:'TAUX PRESSE 2.10%', 0:'Exoneré 0%'}.get(taux,
                      'TVA taux supérieur' if taux is not None else None))
            if taux is not None and taux not in (20, 5.5, 10, 2.1, 0):
                log.append({'gtin': g, 'champ': 'Taux de TVA (604)', 'constat': f"Taux inhabituel : {taux} %."})
        # contacts
        for i in range(min(P.n('contactInformationList.contactNameText'), 3)):
            if P.v(r, 'contactInformationList.contactNameText', i) in (None, ''): continue
            S('185_1', V_CONTACT.get(code(P.v(r, 'contactInformationList.contactTypeCode', i))), i)
            S('185_2', trunc(P.v(r, 'contactInformationList.contactNameText', i), 200), i)
            S('185_4', gtin_txt(P.v(r, 'contactInformationList.contactGLN', i)), i)
            S('2227', trunc(P.v(r, 'contactInformationList.structuredAddressList.streetAddressText', i), 200), i)
            S('2225', trunc(P.v(r, 'contactInformationList.structuredAddressList.postalCodeText', i), 80), i)
            S('2223', trunc(P.v(r, 'contactInformationList.structuredAddressList.cityText', i), 200), i)
            S('2226', trunc(P.v(r, 'contactInformationList.structuredAddressList.provinceStateCodeText', i), 80), i)
            S('2224', V_COUNTRY.get(code(P.v(r, 'contactInformationList.structuredAddressList.countryCode', i))), i)
            S('185_5', 'France', i)
            if i == 0:   # l'extraction ne porte qu'un canal de communication
                S('185_6', V_CHANNEL.get(code(P.v(r, 'contactInformationList.contactCommunicationChannelList.contactCommunicationChannelCode', 0))))
                S('185_7', P.v(r, 'contactInformationList.contactCommunicationChannelList.contactCommunicationChannelValueText', 0))
        prods.append(d)

        # ---------------- logistique ----------------
        SL('2', g); SL('11', chemin); SL('3', P.v(r, 'supplierCode'))
        SL('109', P.v(r, 'brandText')); SL('103', trunc(nom, 200))
        SL('UL2', 'Non'); SL('UL3_1', 'France'); SL('UL3_2', 'GSS Parfumerie')
        ch = chaine(g)
        if ch:
            SL('UL1', V_LOG_TYPE.get(code(L.v(bygtin[ch[-1]], 'typePackaging'))))
        else:
            log.append({'gtin': g, 'champ': 'Hiérarchie logistique', 'constat': "Aucune unité logistique rattachée."})
        cumul = 1
        for b, lug in enumerate(ch):
            lr = bygtin[lug]
            typ = code(L.v(lr, 'typePackaging')); pal = typ == 'PALLET'
            def SB(fid, val):
                if val not in (None, ''): dl[(fid, b)] = val
            SB('PK3', V_LOG_TYPE.get(typ)); SB('PK1', gtin_txt(L.v(lr, 'gtin'))); SB('PK4', 'Finale')
            SB('PK7', trunc(L.v(lr, 'namePublicLong'), 35)); SB('PK8', trunc(L.v(lr, 'namePublicLong'), 200))
            SB('PK9', gtin_txt(L.v(lr, 'children.gtin')))
            q = num(L.v(lr, 'children.quantity')); SB('PK10', q)
            if q: cumul *= q
            SB('PK17', cumul if q else None)
            SB('PK12', L.v(lr, 'quantityOfCompleteLayersContainedInATradeItem'))
            SB('PK16', L.v(lr, 'quantityOfTradeItemsContainedInACompleteLayer'))
            for path, fid in [('isSizedBy.depth','PK20'), ('isSizedBy.width','PK22'), ('isSizedBy.height','PK24')]:
                val, u = L.vu(lr, path); val = num(val)
                if val is not None: SB(fid, round(val/10, 2) if u == 'mm' else val)
            for path, fid in [('grossWeight','PK28'), ('netWeight','PK30')]:
                val, u = L.vu(lr, path); val = num(val)
                if val is not None: SB(fid, round(val/1000, 3) if u == 'g' else val)
            SB('PK32', V_PACK.get(code(L.v(lr, 'packagingInformationList.packagingTypeCode'))))
            SB('PK13', trunc(L.v(lr, 'packagingInformationList.packagingTypeDescriptionTextList'), 200))
            SB('PK35', V_PLATFORM_TC.get(code(L.v(lr, 'packagingInformationList.platformTermsAndConditionsCode'))))
            plat = lab(L.v(lr, 'packagingInformationList.platformTypeCode')) or ''
            for kk, vv in [('1200 x 1000','Palette ISO 2 (120x100 cm)'), ('800 X 1200','Palette ISO 1 (80 x 120 cm)'),
                           ('800 x 1200','Palette ISO 1 (80 x 120 cm)'), ('800 x 600','Palette ISO 0 (60 x 80 cm)'),
                           ('1219 X 1016','Palette 1219 X 1016 mm'), ('sur mesure','Palette sur mesure (Custom platform)')]:
                if kk in plat: SB('PK36', vv); break
            SB('PK37', V_BARCODE.get(code(L.v(lr, 'dataCarrierTypeCode'))))
            SB('PK14', L.v(lr, 'packagingInformationList.packagingLevelNumber'))
            for fid, val in [('PK40','Non'),('PK41','Non'),('PK42','Oui'),('PK43','Non'),('PK44','Non')]:
                SB(fid, val)
            SB('PK46', 'Non' if pal else 'Oui'); SB('PK49', 'Non' if pal else 'Oui')
            SB('PK47', V_UOM_FACT.get(code(L.v(lr, 'orderingUnitOfMeasure'))))
            SB('PK50', L.v(lr, 'orderQuantityMinimum')); SB('PK53', L.v(lr, 'orderQuantityMultiple'))
            SB('PK71', d_fr(L.v(lr, 'startAvailabilityDateTime'))); SB('PK72', d_fr(L.v(lr, 'firstShipDateTime')))
        logis.append(dl)
    wb.close()
    return prods, logis

# ============================ écriture du classeur ============================
TXT = {'2','PK1','PK9','169','2580','172','185_4'}

# Longueurs maximales du dictionnaire Gaia : certaines ne sont pas déclarées dans le
# classeur (le champ 98 par exemple) mais restent contrôlées à l'intégration.
LIMITS = {k: int(v) for k, v in
          json.load(open(_os.path.join(_ICI, 'assets', 'field_limits.json'), encoding='utf-8')).items()}
GLN_FIELDS = {'172', '185_4'}

def cle_gln(g):
    d = str(g)
    if not d.isdigit() or len(d) != 13: return False
    n = [int(x) for x in d]
    s = sum(x * (1 if i % 2 == 0 else 3) for i, x in enumerate(n[:-1]))
    return (10 - s % 10) % 10 == n[-1]

def ecrire(prods, logis, out_path, log, tpl_path=None):
    global TPL
    TPL = tpl_path or TPL
    tpl = load_workbook(TPL)
    import tempfile
    work = tempfile.mkdtemp(prefix='gaia_')      # répertoire temporaire, nettoyé à la fin
    with zipfile.ZipFile(TPL) as z: z.extractall(work)

    n = len(prods); last = FIRST_NEW + n - 1
    def keymap(ws):
        seen = defaultdict(int); out = {}
        for c in range(1, ws.max_column + 1):
            fid = str(ws.cell(row=1, column=c).value)
            lang = ws.cell(row=6, column=c).value
            out[(fid, str(lang) if lang else None, seen[(fid, str(lang) if lang else None)])] = c
            seen[(fid, str(lang) if lang else None)] += 1
        return out

    ss_p = os.path.join(work, 'xl/sharedStrings.xml')
    ss = open(ss_p, encoding='utf-8').read()
    index = {}
    for i, m in enumerate(re.finditer(r'<si>(.*?)</si>', ss, re.S)):
        index.setdefault(re.sub(r'<[^>]+>', '', m.group(1)), i)
    nb_si = len(re.findall(r'<si>', ss)); ajouts = []; nb_str = 0
    def ref(t):
        nonlocal nb_si
        b = str(t)
        if b in index: return index[b]
        index[b] = nb_si
        sp = ' xml:space="preserve"' if b != b.strip() else ''
        ajouts.append(f'<si><t{sp}>{b.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")}</t></si>')
        nb_si += 1
        return index[b]

    total = 0
    for nom, sheetxml, recs in [('Produit', 'xl/worksheets/sheet1.xml', prods),
                                ('Logistique', 'xl/worksheets/sheet2.xml', logis)]:
        kn = keymap(tpl[nom])
        manquants = defaultdict(int)
        cells = {}
        for i, rec in enumerate(recs):
            rn = FIRST_NEW + i
            for (fid, occ), v in rec.items():
                lim = LIMITS.get(fid)
                if lim and isinstance(v, str) and len(v) > lim:
                    court = v[:lim].rsplit(' ', 1)[0] if ' ' in v[:lim] else v[:lim]
                    log.append({'gtin': rec.get(('2', 0), '—'), 'champ': f"{fid} ({nom})",
                                'constat': f"Texte de {len(v)} caractères ramené à {len(court)} (limite {lim})."})
                    v = court
                if fid in GLN_FIELDS and v not in (None, '') and not cle_gln(v):
                    log.append({'gtin': rec.get(('2', 0), '—'), 'champ': f"{fid} ({nom})",
                                'constat': f"GLN « {v} » : clé de contrôle GS1 invalide, champ laissé vide."})
                    continue
                c = (kn.get((fid, 'Français', occ)) or kn.get((fid, None, occ))
                     or kn.get(({'255': '3706'}.get(fid, fid), 'Français', occ))
                     or kn.get(({'255': '3706'}.get(fid, fid), None, occ)))
                if c is None:
                    manquants[(fid, occ)] += 1; continue
                cells[f"{gl(c)}{rn}"] = (v, fid)
        for (fid, occ), k in sorted(manquants.items()):
            lbl = ''
            for cc in range(1, tpl[nom].max_column + 1):
                if str(tpl[nom].cell(row=1, column=cc).value) == fid:
                    lbl = tpl[nom].cell(row=2, column=cc).value; break
            log.append({'gtin': '—', 'champ': f"{nom} / {fid}" + (f" (occurrence {occ+1})" if occ else ""),
                        'constat': f"Aucune colonne dans l'extraction{' pour cette occurrence' if occ else ''}"
                                   f"{' — ' + str(lbl) if lbl else ''}. {k} produit(s)."})
        p = os.path.join(work, sheetxml)
        xml = open(p, encoding='utf-8').read()
        if n > LAST_TPL - FIRST_NEW + 1:
            skel = re.search(r'<row r="%d"[^>]*>.*?</row>' % LAST_TPL, xml, re.S).group(0)
            skel = re.sub(r'(<c r="[A-Z]+)%d"' % LAST_TPL, r'\1{n}"',
                          re.sub(r'<row r="%d"' % LAST_TPL, '<row r="{n}"', skel))
            skel = re.sub(r'<v>.*?</v>', '', skel).replace(' t="s"', '').replace('></c>', '/>')
            xml = xml.replace('</sheetData>', ''.join(skel.format(n=r) for r in range(LAST_TPL + 1, last + 1)) + '</sheetData>')
            xml = re.sub(r'(<dimension ref="A1:[A-Z]+)%d"' % LAST_TPL, r'\g<1>%d"' % last, xml)
            xml = re.sub(r'sqref="([A-Z]+)7:([A-Z]+)%d"' % LAST_TPL, r'sqref="\g<1>7:\g<2>%d"' % last, xml)
        for r_ref, (v, fid) in cells.items():
            if isinstance(v, (int, float)) and not isinstance(v, bool) and fid not in TXT:
                corps, attr = f'<v>{v}</v>', ''
            else:
                corps, attr = f'<v>{ref(v)}</v>', ' t="s"'; nb_str += 1
            xml, k = re.subn(r'<c r="%s"([^>/]*)/>' % r_ref,
                             lambda m: f'<c r="{r_ref}"{m.group(1)}{attr}>{corps}</c>', xml, count=1)
            total += k
        open(p, 'w', encoding='utf-8').write(xml)

    if ajouts:
        m = re.search(r'<sst count="(\d+)" uniqueCount="(\d+)"', ss)
        ss = ss.replace(m.group(0), f'<sst count="{int(m.group(1))+nb_str}" uniqueCount="{int(m.group(2))+len(ajouts)}"', 1)
        ss = ss.replace('</sst>', ''.join(ajouts) + '</sst>')
        open(ss_p, 'w', encoding='utf-8').write(ss)

    with zipfile.ZipFile(TPL) as zin, zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, open(os.path.join(work, info.filename), 'rb').read())
    shutil.rmtree(work)
    return total


# ============================ ligne de commande ============================
def main():
    ap = argparse.ArgumentParser(description="Export SupplierXM -> fichier au format extraction Gaia")
    ap.add_argument('--export', required=True, help="export SupplierXM (.xlsx)")
    ap.add_argument('--template', required=True, help="extraction Gaia servant de gabarit (.xlsx)")
    ap.add_argument('--out', required=True, help="fichier à produire (.xlsx)")
    ap.add_argument('--report', help="journal des anomalies (.json)")
    a = ap.parse_args()
    global TPL
    TPL = a.template
    log = []
    prods, logis = construire(a.export, log)
    n = ecrire(prods, logis, a.out, log, a.template)
    if a.report:
        json.dump(log, open(a.report, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    from collections import Counter
    c = Counter('GLN invalide' if 'clé de contrôle' in x['constat'] else
                ('Texte tronqué' if 'ramené' in x['constat'] else x['champ'].split(' (')[0]) for x in log)
    print(f"{len(prods)} produits | {n} cellules écrites | {len(log)} anomalies")
    for k, v in c.most_common(10):
        print(f"   {v:5} {k}")
    sans = [x for x in log if 'Classification' in x['champ']]
    if sans:
        print("\nATTENTION — catégories sans chemin de classification :")
        for x in sorted({x['constat'] for x in sans}):
            print("   ", x)
        print("   -> compléter CLASSIF / CLASSIF_KIND / GPC avant de livrer le fichier.")

if __name__ == '__main__':
    main()
