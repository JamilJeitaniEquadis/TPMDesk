# -*- coding: utf-8 -*-
"""Génération des fichiers data.xml Equadis pour les visuels et les documents Marionnaud."""
import csv, datetime, json
from collections import defaultdict
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as cif

import argparse, os, datetime as _dt
_ap = argparse.ArgumentParser(description="Export SupplierXM -> data.xml Equadis (visuels et documents)")
_ap.add_argument('--export', required=True)
_ap.add_argument('--outdir', required=True)
_ap.add_argument('--date-defaut', default=_dt.date.today().strftime('%d/%m/%Y'),
                 help="date de début de validité si la source n'en porte pas")
_a = _ap.parse_args()
SRC = _a.export
OUT = _a.outdir.rstrip('/') + '/'
EXPORT_DATE = _a.date_defaut
os.makedirs(OUT, exist_ok=True)

anomalies = []
def note(gtin, champ, msg):
    anomalies.append({'gtin': gtin, 'objet': champ, 'constat': msg})

def code(v):
    if v in (None, ''): return None
    s = str(v); return s.split('∣')[-1].strip() if '∣' in s else s.strip()
def lab(v):
    if v in (None, ''): return None
    s = str(v); return s.split('∣')[0].strip() if '∣' in s else s.strip()
def d_fr(v):
    if v in (None, ''): return None
    s = str(v)[:10]
    try:
        return datetime.datetime.strptime(s, '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return None

# ---------- tables de correspondance ----------
ANGLE_H = {'3': 'ANGLE_CENTRE', '1': 'ANGLE_CENTRE_PLONGEE',
           '2': 'ANGLE_TROIS_QUARTS_DROIT', '4': 'ANGLE_TROIS_QUARTS_GAUCHE'}
ANGLE_V = {'0': 'PARALLELE_SOL', '2': 'VUE_CONTRE_PLONGEE', '1': 'VUE_PLONGEANTE'}
FACE = {'0': 'ANGLE_FACE', '6': 'ANGLE_NON_APPLICABLE', '4': 'ANGLE_COTE_DROIT',
        '2': 'ANGLE_COTE_GAUCHE', '3': 'ANGLE_DESSUS', '5': 'ANGLE_DESSOUS', '7': 'ANGLE_DOS'}
DEFIN = {'0': 'DEFINITION_STANDARD', '1': 'DEFINITION_HAUTE'}
CONTENU = {'1': 'TYPE_CONTENU_EMBALLE_PACKSHOT',   # Produit emballé
           '0': 'TYPE_CONTENU_NU_DEBALLE',          # Produit nu / déballé
           '3': 'TYPE_CONTENU_MISE_EN_SITUATION',   # En situation
           '10': 'TYPE_CONTENU_styled',             # Stylisé
           '5': 'TYPE_CONTENU_EMBALLE_PACKSHOT',    # Optimisé web -> packshot détouré
           '11': 'TYPE_CONTENU_plated',             # Présenté
           '12': 'TYPE_CONTENU_held',               # Pris en main
           '14': 'TYPE_CONTENU_family',             # Avec des produits similaires
           '6': 'TYPE_CONTENU_case',                # En carton
           '8': 'TYPE_CONTENU_brut'}                # Brut / non cuisiné
DOC_TYPE = {'DOC_SAFETY_DATA_SHEET': 'IMAGE_SAFETY_DATA_SHEET',
            'TECHNICAL_DATA_SHEET': 'IMAGE_TECHNICAL_DATA_SHEET',
            'PRODUCT_IMAGE': 'PRODUCT_IMAGE'}

def xml_of(fields):
    out = ['\t<image>']
    for tag, val in fields:
        if val in (None, ''): continue
        if tag == 'descriptionContenu':
            out.append(f'\t\t<descriptionContenu lang="LANG_FRA">{escape(str(val))}</descriptionContenu>')
        else:
            out.append(f'\t\t<{tag}>{escape(str(val))}</{tag}>')
    out.append('\t</image>')
    return '\n'.join(out)

wb = load_workbook(SRC, read_only=True, data_only=True)
S = lambda r, c: r[cif(c) - 1]   # accès par lettre : voir note en fin de fichier
manifest = []

# =====================  VISUELS  =====================
img_rows = list(wb['Images'].iter_rows(min_row=8, values_only=True))
by_gtin = defaultdict(list)
for r in img_rows:
    by_gtin[str(S(r, 'B'))].append(r)

images_xml, images_zip_xml, names = [], [], set()
for gtin, rows in by_gtin.items():
    # le packshot en premier, puis par numéro de séquence
    rows.sort(key=lambda r: (str(S(r, 'D')) != 'vrai', S(r, 'AC') or 99))
    for i, r in enumerate(rows, 1):
        url = str(S(r, 'E'))
        ext = url.rsplit('.', 1)[-1].lower()
        ext = 'jpg' if ext not in ('jpg', 'jpeg', 'png', 'tif', 'tiff') else ext
        fname = f"{gtin}_{i:02d}.{ext}"
        if fname in names: note(gtin, 'fileName', f"Nom de fichier en double : {fname}")
        names.add(fname)
        ext_id = url.rsplit('/', 1)[-1].rsplit('.', 1)[0]

        ah = ANGLE_H.get(code(S(r, 'R')))
        if ah is None:
            ah = 'ANGLE_CENTRE'
            note(gtin, 'horizontalAngle', f"Visuel {i} : angle horizontal absent de la source, valeur ANGLE_CENTRE appliquée.")
        face = FACE.get(code(S(r, 'W')))
        if face is None:
            face = 'ANGLE_NON_APPLICABLE'
            note(gtin, 'mainSideShown', f"Visuel {i} : face absente de la source, valeur ANGLE_NON_APPLICABLE appliquée.")
        dfn = DEFIN.get(code(S(r, 'O')))
        if dfn is None:
            w = S(r, 'AG') or 0
            dfn = 'DEFINITION_HAUTE' if (isinstance(w, (int, float)) and w >= 3000) else 'DEFINITION_STANDARD'
            note(gtin, 'definition', f"Visuel {i} : définition absente de la source, déduite de la largeur ({w} px) -> {dfn}.")
        cont = CONTENU.get(code(S(r, 'K')))
        if cont is None:
            cont = 'TYPE_CONTENU_EMBALLE_PACKSHOT'
            note(gtin, 'contentDescription', f"Visuel {i} : type de contenu absent de la source, packshot par défaut.")
        transparent = 'True' if code(S(r, 'AF')) == '1' or ext == 'png' else 'False'
        ref_type = 'OUT_OF_PACKAGE_IMAGE' if code(S(r, 'K')) == '0' else 'PRODUCT_IMAGE'
        start = d_fr(S(r, 'U')) or EXPORT_DATE

        champs_communs = [
            ('referencedFileTypeCode', ref_type),
            ('horizontalAngle', ah), ('verticalAngle', ANGLE_V.get(code(S(r, 'S')))),
            ('fileEffectiveStartDateTime', start),
            ('fileEffectiveEndDateTime', d_fr(S(r, 'T'))),
            ('canFilesBeEdited', 'False'), ('isFileBackgroundTransparent', transparent),
            ('contentDescription', cont), ('mainSideShown', face), ('definition', dfn),
            ('targetMarket', '250'),
            ('isPrimaryFile', 'True' if str(S(r, 'D')) == 'vrai' else 'False'),
            ('descriptionContenu', 'Image'), ('fileLanguageCode', 'LANG_FRA'),
        ]
        entete = [('gtin', gtin), ('externalID', ext_id), ('action', 'ADD_OR_UPDATE'),
                  ('title', f"Visuel {i}")]
        images_xml.append(xml_of(entete + [('externalUrl', url)] + champs_communs))
        images_zip_xml.append(xml_of(entete + [('fileName', fname)] + champs_communs))
        manifest.append({'type': 'image', 'gtin': gtin, 'url': url, 'fileName': fname,
                         'externalID': ext_id, 'zip': 'images.zip'})

# =====================  DOCUMENTS  =====================
doc_rows = list(wb['Documents'].iter_rows(min_row=8, values_only=True))
by_gtin_d = defaultdict(list)
for r in doc_rows:
    by_gtin_d[str(S(r, 'B'))].append(r)

docs_xml, docs_zip_xml = [], []
for gtin, rows in by_gtin_d.items():
    for i, r in enumerate(rows, 1):
        url = str(S(r, 'D'))
        ext_id = url.rsplit('/', 1)[-1].split('.')[0]
        ref = DOC_TYPE.get(code(S(r, 'F')), 'IMAGE_DOCUMENT')
        suffix = {'IMAGE_SAFETY_DATA_SHEET': 'FDS', 'IMAGE_TECHNICAL_DATA_SHEET': 'FT'}.get(ref, 'DOC')
        is_img = ref == 'PRODUCT_IMAGE'
        fname = f"{gtin}_{suffix}_{i:02d}." + ('jpg' if is_img else 'pdf')
        start = d_fr(S(r, 'H'))
        if not start:
            start = d_fr(S(r, 'E')) or EXPORT_DATE
            note(gtin, 'fileEffectiveStartDateTime', "Document : date de début de validité absente, date de création du fichier utilisée.")
        if is_img:
            note(gtin, 'referencedFileTypeCode', "Ce 'document' est en réalité une image produit : à basculer éventuellement dans le fichier visuels.")
        champs_communs = [('referencedFileTypeCode', ref),
            ('horizontalAngle', 'ANGLE_CENTRE'),
            ('fileEffectiveStartDateTime', start),
            ('fileEffectiveEndDateTime', d_fr(S(r, 'G'))),
            ('canFilesBeEdited', 'False'), ('isFileBackgroundTransparent', 'False'),
            ('contentDescription', 'TYPE_CONTENU_EMBALLE_PACKSHOT'),
            ('mainSideShown', 'ANGLE_NON_APPLICABLE'), ('definition', 'DEFINITION_AUTRE'),
            ('targetMarket', '250'), ('isPrimaryFile', 'False'),
            ('descriptionContenu', 'Fiche de sécurité' if suffix == 'FDS' else 'Document'),
            ('fileLanguageCode', 'LANG_FRA'),
        ]
        entete = [('gtin', gtin), ('externalID', ext_id), ('action', 'ADD_OR_UPDATE'),
                  ('title', f"{'Fiche de sécurité' if suffix == 'FDS' else 'Document'} {i}")]
        docs_xml.append(xml_of(entete + [('externalUrl', url)] + champs_communs))
        docs_zip_xml.append(xml_of(entete + [('fileName', fname)] + champs_communs))
        manifest.append({'type': 'document', 'gtin': gtin, 'url': url, 'fileName': fname,
                         'externalID': ext_id, 'zip': 'documents.zip'})

# =====================  ÉCRITURE  =====================
def write_xml(path, blocks):
    with open(path, 'w', encoding='utf-8', newline='\r\n') as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n<images>\n')
        f.write('\n'.join(blocks))
        f.write('\n</images>\n')

write_xml(OUT + 'data_images.xml', images_xml)                      # hébergement hors Equadis (externalUrl)
write_xml(OUT + 'data_documents.xml', docs_xml)
write_xml(OUT + 'repli_data_images_avec_zip.xml', images_zip_xml)   # repli : hébergement dans Equadis (fileName)
write_xml(OUT + 'repli_data_documents_avec_zip.xml', docs_zip_xml)

with open(OUT + 'Telechargement_visuels.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=['type', 'gtin', 'url', 'fileName', 'externalID', 'zip'], delimiter=';')
    w.writeheader(); w.writerows(manifest)

json.dump(anomalies, open('anomalies_media.json', 'w'), ensure_ascii=False, indent=1)
print('visuels :', len(images_xml), '| documents :', len(docs_xml), '| repli :', len(images_zip_xml), len(docs_zip_xml))
print('manifeste :', len(manifest), '| noms uniques :', len({m['fileName'] for m in manifest}))
print('anomalies :', len(anomalies))

# NOTE : ce script accède encore aux colonnes par leur lettre, contrairement à pipeline.py.
# Les onglets Images et Documents ont montré une structure stable sur les 8 exports traités,
# mais si un export futur décale ces colonnes, basculer sur la classe Feuille de pipeline.py
# (repérage par chemin technique en ligne 4).
