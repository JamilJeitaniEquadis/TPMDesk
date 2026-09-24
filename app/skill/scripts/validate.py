# -*- coding: utf-8 -*-
"""
Contrôle d'un fichier produit, AVANT de le déposer dans Gaia.

    python scripts/validate.py --file sortie.xlsx --template extraction_gaia.xlsx

Trois sources de contrôle, parce qu'aucune ne suffit seule :

  1. Les règles de validation déclarées dans le classeur (listes déroulantes, longueurs).
     -> ne couvrent qu'une partie des champs.
  2. Les longueurs maximales du dictionnaire de champs (assets/field_limits.json).
     -> le classeur ne déclare rien pour certains champs (98, 101...) que la plateforme
        contrôle pourtant. C'est ce trou qui a produit un rejet en production.
  3. La clé de contrôle GS1 sur tous les champs GLN.
     -> la source contient des numéros complétés par des zéros qui ne sont pas des GLN.

Vérifie aussi l'intégrité du fichier : en-têtes intacts et seuls sheet1/sheet2/sharedStrings
modifiés par rapport au gabarit.

Sortie : 0 si tout est propre, 1 sinon.
"""
import argparse, hashlib, json, os, re, sys, zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as gl, column_index_from_string as cif

ICI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIMITS = json.load(open(os.path.join(ICI, 'assets', 'field_limits.json'), encoding='utf-8'))
GLN_FIELDS = {'172', '185_4'}
PREMIERE_LIGNE = 7

def cle_gln(g):
    d = str(g)
    if not d.isdigit() or len(d) != 13: return False
    n = [int(x) for x in d]
    return (10 - sum(x * (1 if i % 2 == 0 else 3) for i, x in enumerate(n[:-1])) % 10) % 10 == n[-1]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', required=True)
    ap.add_argument('--template', required=True)
    a = ap.parse_args()

    tpl = load_workbook(a.template)
    out = load_workbook(a.file)
    listes = {n: tpl[n] for n in ('Lists_Prd', 'Lists_Log') if n in tpl.sheetnames}

    def valeurs(f):
        m = re.match(r"(Lists_\w+)!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", str(f).strip('='))
        if not m or m.group(1) not in listes: return None
        w = listes[m.group(1)]; c = cif(m.group(2))
        return [str(w.cell(row=r, column=c).value)
                for r in range(int(m.group(3)), int(m.group(5)) + 1)
                if w.cell(row=r, column=c).value is not None]

    pb = []
    for feuille in ('Produit', 'Logistique'):
        ws, wt = out[feuille], tpl[feuille]
        ids = {c: str(wt.cell(row=1, column=c).value) for c in range(1, wt.max_column + 1)}
        derniere = ws.max_row
        # 1. règles du classeur
        for dv in wt.data_validations.dataValidation:
            v = valeurs(dv.formula1)
            for rng in str(dv.sqref).split():
                m = re.match(r'([A-Z]+)', rng)
                if not m: continue
                c = cif(m.group(1))
                for r in range(PREMIERE_LIGNE, derniere + 1):
                    val = ws.cell(row=r, column=c).value
                    if val in (None, ''): continue
                    if dv.type == 'list' and v and str(val) not in v:
                        pb.append((feuille, r, gl(c), ids.get(c, ''), 'valeur hors liste', str(val)[:50]))
                    if dv.type == 'textLength' and str(dv.formula1).isdigit() \
                            and len(str(val)) > int(str(dv.formula1)):
                        pb.append((feuille, r, gl(c), ids.get(c, ''), 'longueur (classeur)',
                                   f"{len(str(val))} > {dv.formula1}"))
        # 2. dictionnaire + 3. GLN
        for c in range(1, ws.max_column + 1):
            fid = ids.get(c, ''); lim = LIMITS.get(fid)
            for r in range(PREMIERE_LIGNE, derniere + 1):
                val = ws.cell(row=r, column=c).value
                if val in (None, ''): continue
                if lim and isinstance(val, str) and len(val) > int(lim):
                    pb.append((feuille, r, gl(c), fid, 'longueur (dictionnaire)', f"{len(val)} > {lim}"))
                if fid in GLN_FIELDS and not cle_gln(val):
                    pb.append((feuille, r, gl(c), fid, 'GLN invalide', str(val)))

    # intégrité
    entetes_ok = all(tpl[f].cell(row=r, column=c).value == out[f].cell(row=r, column=c).value
                     for f in ('Produit', 'Logistique')
                     for r in range(1, PREMIERE_LIGNE)
                     for c in range(1, tpl[f].max_column + 1))
    za, zb = zipfile.ZipFile(a.template), zipfile.ZipFile(a.file)
    memes = za.namelist() == zb.namelist()
    modifs = [n for n in za.namelist()
              if hashlib.md5(za.read(n)).hexdigest() != hashlib.md5(zb.read(n)).hexdigest()] if memes else []
    attendus = {'xl/sharedStrings.xml', 'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'}

    lignes = sum(1 for r in range(PREMIERE_LIGNE, out['Produit'].max_row + 1)
                 if out['Produit'].cell(row=r, column=1).value not in (None, ''))
    print(f"Fichier      : {os.path.basename(a.file)}")
    print(f"Produits     : {lignes}")
    print(f"En-têtes     : {'intacts' if entetes_ok else 'MODIFIÉS — à corriger'}")
    print(f"Parties du .xlsx modifiées : {sorted(modifs) if memes else 'liste des fichiers différente'}")
    if memes and set(modifs) - attendus:
        print(f"   ATTENTION : {sorted(set(modifs) - attendus)} ne devrait pas changer.")
        print("   Signe d'un enregistrement par openpyxl au lieu de l'écriture XML directe.")
    print(f"Violations   : {len(pb)}")
    for x in pb[:40]:
        print(f"   {x[0]:11} ligne {x[1]:4} col {x[2]:4} champ {x[3]:7} {x[4]:22} {x[5]}")
    if len(pb) > 40:
        print(f"   ... et {len(pb) - 40} autres")
    ok = not pb and entetes_ok and memes and not (set(modifs) - attendus)
    print("\n" + ("PRÊT À DÉPOSER" if ok else "NE PAS DÉPOSER EN L'ÉTAT"))
    sys.exit(0 if ok else 1)

if __name__ == '__main__':
    main()
