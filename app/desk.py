# -*- coding: utf-8 -*-
"""
Pont entre la page et les scripts du skill, exécutés tels quels dans Pyodide.
Chaque fonction publique renvoie du texte court : c'est ce que Claude lit après un appel d'outil.
"""
import io, os, re, sys, json, contextlib
from collections import Counter

sys.path.insert(0, '/skill/scripts')
import pipeline, inspect_export, validate
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as cif

EXPORT, TEMPLATE, OUT, REPORT = '/work/export.xlsx', '/work/template.xlsx', '/work/output.xlsx', '/work/log.json'
MAX = 6000
DECISIONS = []          # arbitrages pris pendant la session, rejoués à chaque génération
_paths = None           # liste du champ 11 lue dans l'extraction

def _court(t, n=MAX):
    t = t.rstrip()
    return t if len(t) <= n else t[:n] + f"\n… ({len(t) - n} caractères coupés)"

def _run(module, argv):
    buf = io.StringIO(); code = 0
    old = sys.argv
    sys.argv = [module.__name__ + '.py'] + argv
    try:
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
            module.main()
    except SystemExit as e:
        code = e.code if isinstance(e.code, int) else (0 if e.code is None else 1)
    except Exception as e:
        code = 1; buf.write(f"\nERREUR : {type(e).__name__}: {e}")
    finally:
        sys.argv = old
    return code, buf.getvalue()

def _need(*paths):
    manque = [p for p in paths if not os.path.exists(p)]
    if manque:
        noms = {EXPORT: "l'export SupplierXM", TEMPLATE: "l'extraction Gaia (gabarit)", OUT: 'le fichier généré'}
        return 'Manquant : ' + ', '.join(noms.get(p, p) for p in manque) + '. Demander à l\'utilisateur de le déposer.'
    return None

def reset_outputs():
    for p in (OUT, REPORT):
        if os.path.exists(p): os.remove(p)

# ---------------------------------------------------------------- reconnaissance
def inspect():
    e = _need(EXPORT)
    if e: return e
    return _court(_run(inspect_export, ['--export', EXPORT])[1])

def products(category, limit=25):
    """Libellés réels des produits d'une catégorie source (code ou libellé)."""
    e = _need(EXPORT)
    if e: return e
    wb = load_workbook(EXPORT, read_only=True, data_only=True)
    P = pipeline.Feuille(wb['Product'])
    q = str(category).strip().lower(); out = []; n = 0
    for r in P.rows:
        kind = str(P.v(r, 'kind') or '')
        if q and q not in kind.lower() and q != pipeline.code(kind).lower(): continue
        n += 1
        if len(out) < int(limit):
            out.append(' | '.join(str(x) for x in [
                P.v(r, 'gtin'), P.v(r, 'namePublicLong'), pipeline.lab(P.v(r, 'isClassifiedIn')),
                pipeline.lab(P.v(r, 'org1337WebsiteProductCategory')),
                pipeline.code(P.v(r, 'targetConsumerGender'))] if x not in (None, '')))
    wb.close()
    if not n: return f"Aucun produit pour la catégorie « {category} »."
    return _court(f"{n} produit(s) — gtin | libellé | classification source | catégorie web | genre\n" + '\n'.join(out))

# ---------------------------------------------------------------- champ 11
def _liste11():
    global _paths
    if _paths is not None: return _paths
    wb = load_workbook(TEMPLATE)
    ws = wb['Produit']
    col = next(c for c in range(1, ws.max_column + 1) if str(ws.cell(row=1, column=c).value) == '11')
    lettre = ws.cell(row=1, column=col).column_letter
    vals = []
    for dv in ws.data_validations.dataValidation:
        cols = {re.match(r'([A-Z]+)', x).group(1) for x in str(dv.sqref).split()}
        if lettre not in cols: continue
        m = re.match(r"(Lists_\w+)!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", str(dv.formula1).strip('='))
        if not m: continue
        w = wb[m.group(1)]; c = cif(m.group(2))
        vals = [str(w.cell(row=r, column=c).value) for r in range(int(m.group(3)), int(m.group(5)) + 1)
                if w.cell(row=r, column=c).value is not None]
        break
    _paths = vals
    return vals

def search_paths(query, limit=40):
    e = _need(TEMPLATE)
    if e: return e
    vals = _liste11()
    if not vals: return "L'extraction ne déclare pas de liste pour le champ 11."
    mots = [m for m in str(query).lower().split() if m]
    hits = [v for v in vals if all(m in v.lower() for m in mots)]
    if not hits and mots:
        hits = [v for v in vals if any(m in v.lower() for m in mots)]
    if not hits: return f"Aucun chemin ne contient « {query} » ({len(vals)} chemins dans la liste)."
    return _court(f"{len(hits)} chemin(s) sur {len(vals)} (copier au caractère près) :\n" +
                  '\n'.join(hits[:int(limit)]))

def set_classification(category, path, reason=''):
    e = _need(TEMPLATE)
    if e: return e
    vals = _liste11()
    if vals and path not in vals:
        proches = [v for v in vals if v.split('|')[-1].strip() == str(path).split('|')[-1].strip()]
        return ("REFUSÉ : ce chemin n'est pas dans la liste du champ 11, il serait rejeté. "
                + ("Entrées avec la même brique :\n" + '\n'.join(proches[:10]) if proches else
                   "Chercher avec search_classification_paths et copier l'entrée exacte."))
    k = pipeline.code(category)
    pipeline.CLASSIF_KIND[k] = path
    DECISIONS[:] = [d for d in DECISIONS if not (d['type'] == 'classification' and d['key'] == k)]
    DECISIONS.append({'type': 'classification', 'key': k, 'value': path, 'reason': reason})
    reset_outputs()
    return f"Catégorie {k} → {path}. Enregistré comme arbitrage à faire valider par le client."

TABLES = {n: getattr(pipeline, n) for n in dir(pipeline) if n.startswith('V_') and isinstance(getattr(pipeline, n), dict)}

def set_value(table, source, target, reason=''):
    t = TABLES.get(str(table))
    if t is None: return f"Table inconnue. Tables : {', '.join(sorted(TABLES))}"
    t[str(source)] = target
    DECISIONS[:] = [d for d in DECISIONS if not (d['type'] == table and d['key'] == source)]
    DECISIONS.append({'type': table, 'key': str(source), 'value': target, 'reason': reason})
    reset_outputs()
    return f"{table}[{source!r}] = {target!r}."

def decisions_json():
    return json.dumps(DECISIONS, ensure_ascii=False)

def restore(js):
    for d in json.loads(js or '[]'):
        if d['type'] == 'classification': pipeline.CLASSIF_KIND[d['key']] = d['value']
        elif d['type'] in TABLES: TABLES[d['type']][d['key']] = d['value']
        DECISIONS.append(d)

# ---------------------------------------------------------------- génération et contrôle
def generate():
    e = _need(EXPORT, TEMPLATE)
    if e: return e
    reset_outputs()
    global _paths
    code, txt = _run(pipeline, ['--export', EXPORT, '--template', TEMPLATE, '--out', OUT, '--report', REPORT])
    if code or not os.path.exists(OUT):
        return _court('ÉCHEC de la génération.\n' + txt)
    return _court(txt + ("\n\nArbitrages appliqués : " + str(len(DECISIONS)) if DECISIONS else ''))

def check():
    e = _need(OUT, TEMPLATE)
    if e: return e
    code, txt = _run(validate, ['--file', OUT, '--template', TEMPLATE])
    return _court(('CODE RETOUR 0 — déposable\n' if code == 0 else f'CODE RETOUR {code} — non déposable\n') + txt)

def anomalies(limit=60):
    if not os.path.exists(REPORT): return 'Pas de journal : lancer la génération.'
    log = json.load(open(REPORT, encoding='utf-8'))
    c = Counter(x['champ'] for x in log)
    lignes = [f"{len(log)} anomalies. Par champ :"] + [f"  {v:4}  {k}" for k, v in c.most_common(20)]
    lignes += ['', 'Détail :'] + [f"  {x['gtin']} | {x['champ']} | {x['constat']}" for x in log[:int(limit)]]
    return _court('\n'.join(lignes))

def read_reference(name):
    n = re.sub(r'[^a-z\-]', '', str(name).lower())
    p = f'/skill/references/{n}.md'
    if not os.path.exists(p):
        return 'Références : ' + ', '.join(sorted(f[:-3] for f in os.listdir('/skill/references')))
    return _court(open(p, encoding='utf-8').read(), 9000)

def new_template():
    global _paths
    _paths = None
    reset_outputs()
