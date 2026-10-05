# -*- coding: utf-8 -*-
"""
Pont de l'onglet « Excel → Equa XML » : export SupplierXM (ou extraction Gaia) -> un XML par GLN, contrôlé.
Chaîne : pipeline.py (export -> extraction Gaia) puis excel_to_xml.py (extraction -> XML) puis check_xml.py.
"""
import io, os, re, sys, json, glob, shutil, contextlib, warnings
warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); sys.path.insert(0, '/skill/scripts')
from openpyxl import load_workbook
import excel_to_xml as X, check_xml

def _run(main, argv):
    buf, old = io.StringIO(), sys.argv
    sys.argv = ['x'] + argv
    try:
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf): main()
        return True, buf.getvalue()
    except SystemExit as e:
        return (e.code in (0, None)), buf.getvalue()
    except Exception as e:
        return False, buf.getvalue() + f"\nERREUR : {type(e).__name__}: {e}"
    finally:
        sys.argv = old

def nature(path):
    noms = load_workbook(path, read_only=True).sheetnames
    if 'Product' in noms and 'Logistical units' in noms: return 'export'
    if 'Produit' in noms: return 'extraction'
    return 'inconnu : ' + ', '.join(noms[:6])

def convertir(src, nom, template, gln_lignes='', alias_json='{}', client='MARIONNAUD', sortie='/work/xml'):
    """Renvoie un JSON : nature, fichiers (contrôlés), rapport, journal, chemin de l'extraction."""
    shutil.rmtree(sortie, ignore_errors=True); os.makedirs(sortie)
    base = os.path.splitext(nom)[0]
    m = re.match(r'export_\d+_\d{8}(\d{6})$', base)                 # export_28195_20260630112352 -> lot 112352
    lot = 'Extraction_GAIA_' + (client.upper() + '_' if client else '') + 'export_' + (m.group(1) if m else base)
    out = {'nature': nature(src), 'fichiers': [], 'rapport': '', 'journal': '', 'extraction': None}
    export = None
    if out['nature'] == 'export':
        import pipeline
        extraction = os.path.join(sortie, lot + '.xlsx')
        ok, txt = _run(pipeline.main, ['--export', src, '--template', template, '--out', extraction,
                                       '--report', os.path.join(sortie, 'anomalies.json')])
        out['journal'] += txt
        if not ok: out['erreur'] = 'pipeline.py a échoué'; return json.dumps(out, ensure_ascii=False)
        out['extraction'], export = extraction, src
    elif out['nature'] == 'extraction':
        extraction = os.path.join(sortie, (base if base.startswith('Extraction') else lot) + '.xlsx'); shutil.copy(src, extraction)
    else:
        out['erreur'] = 'Fichier non reconnu (' + out['nature'] + ') : déposer un export SupplierXM ou une extraction Gaia.'
        return json.dumps(out, ensure_ascii=False)
    X.CORRECTIONS.clear(); del X.A_CORRIGER[:]; X.NOMS_PARTIES.clear()
    try: X.ALIAS = {k: v for k, v in json.loads(alias_json or '{}').items() if not k.startswith('_')}
    except ValueError: out['erreur'] = 'Alias GLN : JSON invalide'; return json.dumps(out, ensure_ascii=False)
    argv = ['--excel', extraction, '--outdir', os.path.join(sortie, 'xml')]
    if export: argv += ['--export', export]
    for l in (gln_lignes or '').splitlines():
        if '=' in l and l.strip(): argv += ['--gln', l.strip()]
    ok, txt = _run(X.main, argv)
    out['journal'] += '\n' + txt
    if not ok: out['erreur'] = 'excel_to_xml.py a échoué'; return json.dumps(out, ensure_ascii=False)
    rap = glob.glob(os.path.join(sortie, 'xml', '*_rapport.txt'))
    out['rapport'] = open(rap[0], encoding='utf-8').read() if rap else ''
    out['rapport_chemin'] = rap[0] if rap else None
    for r in check_xml.verifier(os.path.join(sortie, 'xml')):
        r['chemin'] = os.path.join(sortie, 'xml', r['fichier'])
        out['fichiers'].append(r)
    out['fichiers'].sort(key=lambda r: -r['produits'])
    out['regles'] = check_xml.REGLES
    out['alias_defaut'] = {k: v for k, v in json.load(open(os.path.join(HERE, 'gln_alias.json'))).items() if not k.startswith('_')}
    return json.dumps(out, ensure_ascii=False)

def alias_defaut():
    return json.dumps({k: v for k, v in json.load(open(os.path.join(HERE, 'gln_alias.json'))).items() if not k.startswith('_')}, indent=1)
